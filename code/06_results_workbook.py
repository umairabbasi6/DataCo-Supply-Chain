"""Bundle every results/*.csv table into one Excel workbook (results/ALL_RESULTS.xlsx).
Sheets are ordered A1..A3, B1..B8, C0..C13 and an Index sheet comes first."""
import glob, os, re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from common import RES

TITLES = {
    "A1": "DataCo leakage tests", "A2": "DataCo inflation decomposition (XGBoost)",
    "A3": "RF 0.888 vs 0.9608 explained", "B1": "DS2 randomness / association tests",
    "B2": "DS2 learnability under 6 label rules", "B3": "DS2 permutation test",
    "B4": "DS3 label determinism", "B5": "Fused corpus ablation",
    "B6": "Recall-threshold check", "B7": "K-Means clusters vs source",
    "B8": "Cluster metrics", "C0": "DataCo feature set (kept / excluded)",
    "C1": "Benchmark per fold", "C2": "Benchmark mean ± SD (grouped 5-fold)",
    "C3": "Pooled OOF confusion matrices", "C4": "DeLong tests",
    "C5": "Temporal hold-out", "C6": "Permutation importance (XGBoost)",
    "C7": "XGBoost sensitivity grid", "C8": "Before vs after (paper vs corrected)",
    "C9": "MLP selection grid", "C10": "Late rate by shipping mode x order hour",
    "C11": "Lookup-rule ceiling (in-sample and out-of-fold) vs ML",
    "C12": "Late rate by payment Type x shipping mode",
    "C13": "Lookup vs ML bootstrap (order-level)",
}

def sort_key(path):
    code = os.path.basename(path).split("_")[0]
    m = re.match(r"([A-Z])(\d+)", code)
    return (m.group(1), int(m.group(2))) if m else ("Z", 0)

files = sorted((f for f in glob.glob(f"{RES}/*.csv") if "permutation_null" not in f),
               key=sort_key)
codes = [os.path.basename(f).split("_")[0] for f in files]
dups = {c for c in codes if codes.count(c) > 1}
if dups:
    raise SystemExit(f"Duplicate table codes in results/: {sorted(dups)} "
                     "(remove stray copies such as *_1.csv)")
out = f"{RES}/ALL_RESULTS.xlsx"
head = PatternFill("solid", fgColor="2A78D6")
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    idx = []
    for f, code in zip(files, codes):
        df = pd.read_csv(f)
        df.to_excel(xw, sheet_name=code, index=False)
        idx.append((code, TITLES.get(code, ""), os.path.basename(f), len(df)))
    pd.DataFrame(idx, columns=["Sheet", "Table", "CSV file", "Rows"]).to_excel(
        xw, sheet_name="Index", index=False)
    for ws in xw.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, col in enumerate(ws.columns, 1):
            w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 70)
        ws.freeze_panes = "A2"
    xw.book.move_sheet("Index", offset=-len(xw.book.worksheets) + 1)
print(f"Saved {out} ({len(files)} tables)")